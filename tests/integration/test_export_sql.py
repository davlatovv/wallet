"""Export use case against the real repository (SQLite)."""
import csv
import io
from datetime import date, datetime, timezone
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app.application.use_cases.export.export_transactions import (
    ExportFormat, ExportTransactionsUseCase,
)
from app.domain.entities.transaction import TransactionType
from app.domain.exceptions.base import ValidationError
from app.infrastructure.db.repositories.transaction import SQLAlchemyTransactionRepository

D = Decimal


@pytest.fixture
def repo(session):
    return SQLAlchemyTransactionRepository(session)


async def _seed(repo, year, month, day=15):
    tx = await repo.create(1, D("500"), TransactionType.INCOME, None, None)
    await repo.update(tx.id, 1, {"created_at": datetime(year, month, day, tzinfo=timezone.utc)})
    tx2 = await repo.create(1, D("120"), TransactionType.EXPENSE, None, "coffee")
    await repo.update(tx2.id, 1, {"created_at": datetime(year, month, day, tzinfo=timezone.utc)})


async def test_csv_export_for_a_specific_month(repo):
    await _seed(repo, 2026, 3)
    await _seed(repo, 2026, 4)  # different month, must be excluded
    result = await ExportTransactionsUseCase(repo).execute(1, ExportFormat.CSV, date(2026, 3, 1))
    assert result.content_type == "text/csv" and result.filename == "transactions_2026-03.csv"
    rows = list(csv.reader(io.StringIO(result.content.decode("utf-8-sig"))))
    assert len(rows) == 3  # header + 2 rows
    assert {rows[1][1], rows[2][1]} == {"Доход", "Расход"}


async def test_xlsx_export_for_a_specific_month(repo):
    await _seed(repo, 2026, 5)
    result = await ExportTransactionsUseCase(repo).execute(1, ExportFormat.XLSX, date(2026, 5, 20))
    assert result.content_type.endswith("spreadsheetml.sheet")
    assert result.filename == "transactions_2026-05.xlsx"
    ws = load_workbook(io.BytesIO(result.content)).active
    assert ws.cell(row=2, column=3).value == 500.0 or ws.cell(row=3, column=3).value == 500.0


async def test_defaults_to_current_month(repo):
    now = datetime.now(timezone.utc)
    await _seed(repo, now.year, now.month, day=min(now.day, 27))
    result = await ExportTransactionsUseCase(repo).execute(1, ExportFormat.CSV)
    assert result.filename == f"transactions_{now.year:04d}-{now.month:02d}.csv"


async def test_isolated_per_user(repo):
    await _seed(repo, 2026, 6)
    result = await ExportTransactionsUseCase(repo).execute(2, ExportFormat.CSV, date(2026, 6, 1))
    rows = list(csv.reader(io.StringIO(result.content.decode("utf-8-sig"))))
    assert len(rows) == 1  # header only


@pytest.mark.parametrize("bad_month", [date(1999, 1, 1), date(2101, 1, 1)])
async def test_out_of_range_year_rejected(repo, bad_month):
    with pytest.raises(ValidationError):
        await ExportTransactionsUseCase(repo).execute(1, ExportFormat.CSV, bad_month)
